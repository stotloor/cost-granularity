"""Load the Cost Granularity workbook into the cost_granularity Postgres schema.

    python load_workbook.py path/to/Cost_Granularity_Sample_Data_v2.xlsx \
        [--dsn postgresql://localhost/cost_granularity] [--fiscal-year 2026]

The DSN falls back to $DATABASE_URL, then to libpq defaults (PGHOST, PGDATABASE, ...).
Run schema.sql first. The load is a single transaction: it replaces all business
data (etl_load history is kept), so re-running is safe.

Only input cells are loaded; formula columns are recomputed by views. Sheet
headers are a fixed contract and the functional dependencies the schema relies
on are verified, so a changed workbook fails loudly instead of loading wrong.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import sys
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import openpyxl
import psycopg

SAP_HEADERS = ['YTD TOTAL', 'In Copperleaf?', '.', 'Company Code', 'Project Definition',
               'Project Definition', 'WBS Element', 'WBS Element', 'Object Class',
               'Person Responsible', 'Initiative ID', 'Scheme No/Proj Ref',
               'WBS Regulatory Activity', 'WBS Regulatory Activity | Posting period',
               *[str(p) for p in range(1, 13)]]
COPPERLEAF_HEADERS = ['InvestmentCode', 'InvestmentName', 'Area', 'Facility',
                      'FinalSubmissionInvIdT3', 'FinalSubmissionVolumeT3', 'InvestmentVolume',
                      'InvIdCode', 'PriceBaseOfEnteredCost', 'Product', 'Site', 'SiteId',
                      'FundingPositionT3Proposed', 'InterventionDriver', 'SAPProject',
                      'PercentageOfWBS', 'IDTTotalECC', 'PACCode']
PO_HEADERS = ['PO Number', 'Supplier ID', 'Contract Ref', 'PO Description', 'PO Date',
              'Currency', 'Posting Level', 'Summary Cost Object', 'PO Value (£)',
              'Actual Posted (£)', 'Allocation Basis', 'Status', 'Total Allocation %',
              'Allocated Actual (£)', 'Unallocated Difference (£)']
ALLOC_HEADERS = ['Allocation ID', 'PO Number', 'Project Definition', 'Project Name',
                 'WBS Element', 'WBS Description', 'Investment ID', 'Allocation Rule',
                 'Allocation %', 'PO Actual (£)', 'Calculated Allocation (£)',
                 'Manual Adjustment (£)', 'Final Allocated Actual (£)',
                 'Allocation Rationale', 'Review Status']
EXCEPTION_HEADERS = ['Check ID', 'Check', 'Sheet', 'Where', 'What was planted',
                     'Requirement', 'Relevant idea']

# Dummy / pending regulatory activity codes (Seeded Exceptions C1).
PLACEHOLDER_ACTIVITY_CODES = {'W99', 'QQQ'}

TABLES = ['seeded_exception', 'po_allocation', 'purchase_order', 'contract', 'supplier',
          'copperleaf_investment', 'site', 'sap_wbs_actual', 'wbs_element', 'sap_project',
          'regulatory_activity', 'allocation_basis', 'fiscal_period']


class WorkbookError(ValueError):
    pass


def text(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def money(v):
    return None if v is None else Decimal(str(v)).quantize(Decimal('0.01'))


def decimal(v):
    return None if v is None else Decimal(str(v))


def check_headers(sheet, row, expected, first_col=1):
    actual = [text(c.value) for c in sheet[row][first_col - 1:first_col - 1 + len(expected)]]
    if actual != expected:
        diff = [(i + first_col, e, a) for i, (e, a) in enumerate(zip(expected, actual)) if e != a]
        raise WorkbookError(f"{sheet.title!r} row {row} headers changed (col, expected, actual): {diff}")


def data_rows(sheet, start_row, first_col, width, key_prefix=None):
    for r in sheet.iter_rows(min_row=start_row, min_col=first_col,
                             max_col=first_col + width - 1, values_only=True):
        key = text(r[0])
        if key is None or (key_prefix and not key.startswith(key_prefix)):
            continue  # blank rows and the PO Summary 'Totals' row
        yield [text(v) if isinstance(v, str) else v for v in r]


def single(groups, what):
    """Collapse {key: set(values)} to {key: value}; fail if a key has several values."""
    bad = {k: v for k, v in groups.items() if len(v) != 1}
    if bad:
        raise WorkbookError(f"{what} is not unique per key: {bad}")
    return {k: next(iter(v)) for k, v in groups.items()}


def unique(rows, idx, what):
    seen = set()
    for r in rows:
        if r[idx] in seen:
            raise WorkbookError(f"duplicate {what}: {r[idx]}")
        seen.add(r[idx])


def fiscal_periods(fy):
    out = []
    for p in range(1, 13):
        y, m = (fy, p + 3) if p <= 9 else (fy + 1, p - 9)
        start = dt.date(y, m, 1)
        end = (dt.date(y + (m == 12), m % 12 + 1, 1)) - dt.timedelta(days=1)
        out.append((fy, p, start, end))
    return out


def extract(path: Path, fy: int):
    wb = openpyxl.load_workbook(path, data_only=False)

    # ---- SAP Data (header row 2, columns B:AA) ----
    ws = wb['SAP Data']
    check_headers(ws, 2, SAP_HEADERS, first_col=2)
    sap = list(data_rows(ws, 3, 5, 23))  # from Company Code (E)
    unique(sap, 3, 'WBS element')
    proj = defaultdict(lambda: defaultdict(set))
    activity = defaultdict(set)
    for (cc, pdef, pname, wbs, wdesc, ocls, person, ini, scheme, act, actdesc, *_p) in sap:
        for k, v in (('name', pname), ('cc', cc), ('person', person), ('ini', ini), ('scheme', scheme)):
            proj[k][pdef].add(v)
        activity[act].add(actdesc)
    names, ccs, persons, inis, schemes = (single(proj[k], f'SAP project {k}')
                                          for k in ('name', 'cc', 'person', 'ini', 'scheme'))
    activity = single(activity, 'regulatory activity description')

    out = {
        'fiscal_period': fiscal_periods(fy),
        'regulatory_activity': [(c, d, c in PLACEHOLDER_ACTIVITY_CODES) for c, d in sorted(activity.items())],
        'sap_project': [(p, names[p], ccs[p], persons[p], inis[p], schemes[p]) for p in sorted(names)],
        'wbs_element': [(r[3], r[1], r[4], r[5], r[9]) for r in sap],
        'sap_wbs_actual': [(r[3], fy, i + 1, money(v))
                           for r in sap for i, v in enumerate(r[11:23]) if v is not None],
    }
    for r in sap:
        for v in r[11:23]:
            if v is not None and not isinstance(v, (int, float)):
                raise WorkbookError(f"non-numeric SAP amount on {r[3]}: {v!r}")

    # ---- Copperleaf Data (header row 1) ----
    ws = wb['Copperleaf Data']
    check_headers(ws, 1, COPPERLEAF_HEADERS)
    cl = list(data_rows(ws, 2, 1, 18))
    unique(cl, 0, 'investment code')
    site = defaultdict(set)
    for r in cl:
        if r[3] != r[10]:
            raise WorkbookError(f"{r[0]}: Facility {r[3]!r} differs from Site {r[10]!r}")
        site[r[11]].add((r[10], r[2]))
    site = single(site, 'site name/area')
    out['site'] = [(sid, n, a) for sid, (n, a) in sorted(site.items())]
    out['copperleaf_investment'] = [
        (r[0], r[1], r[11], r[4], r[5], r[6], r[7], r[8], r[9], r[12], r[13], r[14],
         decimal(r[15]), money(r[16]), r[17]) for r in cl]

    # ---- PO Summary (header row 4) ----
    ws = wb['PO Summary']
    check_headers(ws, 4, PO_HEADERS)
    po = list(data_rows(ws, 5, 1, 15, key_prefix='PO-'))
    unique(po, 0, 'PO number')
    contract = single({c: {s for _, s, c2, *_ in po if c2 == c} for c in {r[2] for r in po}},
                      'contract supplier')
    out['supplier'] = [(s,) for s in sorted({r[1] for r in po})]
    out['contract'] = sorted(contract.items())
    po_basis = {r[0]: r[10] for r in po}
    out['purchase_order'] = [
        (r[0], r[1], r[2], r[3], r[4].date() if isinstance(r[4], dt.datetime) else r[4], r[5], r[6],
         None if r[7] == 'MULTI-PROJECT' else r[7], money(r[8]), money(r[9]), r[10], r[11])
        for r in po]

    # ---- Actual Allocation (header row 4) ----
    ws = wb['Actual Allocation']
    check_headers(ws, 4, ALLOC_HEADERS)
    al = list(data_rows(ws, 5, 1, 15))
    unique(al, 0, 'allocation ID')
    for r in al:
        if r[2] in names and r[3] != names[r[2]]:
            raise WorkbookError(f"{r[0]}: Project Name {r[3]!r} != SAP name {names[r[2]]!r}")
        if r[1] in po_basis and r[7] != po_basis[r[1]]:
            raise WorkbookError(f"{r[0]}: Allocation Rule {r[7]!r} != PO basis {po_basis[r[1]]!r}")
    out['po_allocation'] = [(r[0], r[1], r[2], r[4], r[5], r[6], decimal(r[8]),
                             money(r[11] or 0), r[13], r[14]) for r in al]
    out['allocation_basis'] = [(b,) for b in sorted(set(po_basis.values()) | {r[7] for r in al})]

    # ---- Seeded Exceptions (header row 4) ----
    ws = wb['Seeded Exceptions']
    check_headers(ws, 4, EXCEPTION_HEADERS)
    out['seeded_exception'] = [tuple(r) for r in data_rows(ws, 5, 1, 7)]
    return out


INSERTS = {
    'fiscal_period': 'fiscal_year, period, period_start, period_end',
    'regulatory_activity': 'activity_code, description, is_placeholder',
    'allocation_basis': 'allocation_basis',
    'site': 'site_id, site_name, area',
    'supplier': 'supplier_id',
    'contract': 'contract_ref, supplier_id',
    'sap_project': 'project_definition, project_name, company_code, person_responsible, initiative_id, scheme_ref',
    'wbs_element': 'wbs_element, project_definition, description, object_class, activity_code',
    'sap_wbs_actual': 'wbs_element, fiscal_year, period, amount',
    'copperleaf_investment': ('investment_code, investment_name, site_id, final_submission_inv_id_t3, '
                              'final_submission_volume_t3, investment_volume, inv_id_code, price_base, '
                              'product, funding_position_t3, intervention_driver, sap_project, '
                              'percentage_of_wbs, idt_total_ecc, pac_code'),
    'purchase_order': ('po_number, supplier_id, contract_ref, description, po_date, currency, posting_level, '
                       'summary_project_definition, po_value, actual_posted, allocation_basis, status'),
    'po_allocation': ('allocation_id, po_number, project_definition, wbs_element, wbs_description, '
                      'initiative_id, allocation_pct, manual_adjustment, allocation_rationale, review_status'),
    'seeded_exception': 'check_id, check_name, source_sheet, location, what_was_planted, requirement, relevant_idea',
}


def load(dsn: str, path: Path, data: dict) -> dict:
    counts = {t: len(data[t]) for t in INSERTS}
    with psycopg.connect(dsn) as conn, conn.cursor() as cur:  # one transaction
        cur.execute('SET search_path = cost_granularity')
        cur.execute(f"TRUNCATE {', '.join(TABLES)}")
        for table, cols in INSERTS.items():  # parent tables first
            ph = ', '.join(['%s'] * len(cols.split(',')))
            cur.executemany(f'INSERT INTO {table} ({cols}) VALUES ({ph})', data[table])
        cur.execute('INSERT INTO etl_load (source_file, source_sha256, row_counts) VALUES (%s, %s, %s)',
                    (path.name, hashlib.sha256(path.read_bytes()).hexdigest(), json.dumps(counts)))
    return counts


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('workbook', type=Path)
    ap.add_argument('--dsn', default=os.environ.get('DATABASE_URL', ''))
    ap.add_argument('--fiscal-year', type=int, default=2026,
                    help='Starting calendar year of the fiscal year the SAP periods belong to (default 2026 = FY26/27)')
    args = ap.parse_args(argv)
    try:
        data = extract(args.workbook, args.fiscal_year)
    except WorkbookError as e:
        sys.exit(f'workbook rejected: {e}')
    counts = load(args.dsn, args.workbook, data)
    width = max(map(len, counts))
    for t, n in counts.items():
        print(f'{t:<{width}}  {n:>4}')


if __name__ == '__main__':
    main()
